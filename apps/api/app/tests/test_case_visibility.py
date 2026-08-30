import io
import uuid

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete, select

from app.database.session import SessionLocal
from app.main import app
from app.modules.access.models import AccessLevelAssignment
from app.modules.api import password_hash
from app.modules.models import Case, Environment, EnvironmentMembership, RequestType, User

client = TestClient(app)


def login(email: str, password: str) -> dict[str, str]:
    response = client.post("/api/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def test_environment_manager_visibility_is_shared_by_every_case_consumer() -> None:
    marker = uuid.uuid4().hex[:8]
    password = "Manager123!"
    with SessionLocal() as db:
        admin = db.scalar(select(User).where(User.is_system_admin.is_(True)))
        environment_rows = list(db.execute(
            select(Environment, RequestType).join(
                RequestType, RequestType.environment_id == Environment.id
            ).where(Environment.is_active.is_(True)).limit(1)
        ))
        assert admin and len(environment_rows) == 1
        managed_environment, managed_type = environment_rows[0]
        other_environment = Environment(code=f"VIS-{marker}", name_he="סביבה אחרת",
            name_en="Other environment", description=None, is_active=True)
        db.add(other_environment);db.flush()
        other_type = RequestType(environment_id=other_environment.id, code=f"TYPE-{marker}",
            name_he="סוג אחר", name_en="Other type", description=None, is_active=True,
            default_priority_id=None, default_sub_priority_id=None,
            default_assignee_user_id=None, default_assignee_group_id=None,
            workflow_definition_id=None)
        db.add(other_type);db.flush()
        manager = User(email=f"uri.manager.{marker}@example.com", display_name="אורי לוי",
                       password_hash=password_hash.hash(password), is_active=True, status="active")
        ordinary = User(email=f"ordinary.{marker}@example.com", display_name="משתמש רגיל",
                        password_hash=password_hash.hash(password), is_active=True, status="active")
        db.add_all([manager, ordinary]); db.flush()
        db.add(EnvironmentMembership(environment_id=managed_environment.id, user_id=manager.id,
            role_id=None, source="manual", is_active=True, is_environment_manager=True))
        db.add(AccessLevelAssignment(domain_code="report_cases", user_id=manager.id, group_id=None,
            environment_id=managed_environment.id, access_level="view", created_by=admin.id))
        managed_case = Case(case_number=f"CASE-VIS-{marker}", environment_id=managed_environment.id,
            request_type_id=managed_type.id, reporter_id=admin.id, requester_id=admin.id,
            assignee_id=None, title="CASE-000021 equivalent", description="manager visibility",
            is_locked=True, locked_by=admin.id, lock_reason="visibility regression")
        other_case = Case(case_number=f"CASE-OTHER-{marker}", environment_id=other_environment.id,
            request_type_id=other_type.id, reporter_id=admin.id, requester_id=admin.id,
            assignee_id=None, title="Environment B", description="must remain hidden")
        db.add_all([managed_case, other_case]); db.commit()
        manager_id, managed_case_id = manager.id, managed_case.id
        managed_number, other_number = managed_case.case_number, other_case.case_number
        managed_environment_id = managed_environment.id

    manager_headers = login(f"uri.manager.{marker}@example.com", password)
    listed = client.get("/api/cases", headers=manager_headers)
    assert listed.status_code == 200
    assert managed_number in {row["case_number"] for row in listed.json()}
    assert other_number not in {row["case_number"] for row in listed.json()}

    workspace = client.get("/api/cases/workspace/query?activity_state=all", headers=manager_headers)
    assert workspace.status_code == 200
    assert managed_number in {row["case_number"] for row in workspace.json()["items"]}
    assert other_number not in {row["case_number"] for row in workspace.json()["items"]}

    unfiltered_report = client.get("/api/reports/cases", headers=manager_headers)
    assert unfiltered_report.status_code == 200
    assert managed_number in {row["case_number"] for row in unfiltered_report.json()["items"]}
    assert other_number not in {row["case_number"] for row in unfiltered_report.json()["items"]}
    filtered_report = client.get(
        f"/api/reports/cases?environment_id={managed_environment_id}&search={managed_number}",
        headers=manager_headers,
    )
    assert filtered_report.status_code == 200
    assert [row["case_number"] for row in filtered_report.json()["items"]] == [managed_number]

    direct = client.get(f"/api/cases/{managed_case_id}", headers=manager_headers)
    assert direct.status_code == 200 and direct.json()["is_locked"] is True
    assert direct.json()["permissions"]["can_edit"] is True
    edited = client.patch(f"/api/cases/{managed_case_id}", headers=manager_headers,
        json={"title": "נערך בידי מנהל הסביבה", "version": direct.json()["version"]})
    assert edited.status_code == 200, edited.text

    exported = client.get(
        f"/api/reports/cases/export?environment_id={managed_environment_id}&search={managed_number}",
        headers=manager_headers,
    )
    assert exported.status_code == 200
    sheet = load_workbook(io.BytesIO(exported.content), read_only=True)["קריאות"]
    assert managed_number in {str(row[0]) for row in sheet.iter_rows(min_row=2, values_only=True)}

    admin_headers = login("admin@example.com", "Admin123!")
    admin_rows = client.get("/api/cases", headers=admin_headers).json()
    assert {managed_number, other_number} <= {row["case_number"] for row in admin_rows}
    impersonated = client.post("/api/impersonation/start", headers=admin_headers,
                               json={"user_id": str(manager_id)})
    assert impersonated.status_code == 200
    impersonated_headers = {"Authorization": f"Bearer {impersonated.json()['access_token']}"}
    assert client.get(f"/api/cases/{managed_case_id}", headers=impersonated_headers).status_code == 200
    impersonated_rows = client.get("/api/cases", headers=impersonated_headers).json()
    assert managed_number in {row["case_number"] for row in impersonated_rows}
    assert other_number not in {row["case_number"] for row in impersonated_rows}

    ordinary_headers = login(f"ordinary.{marker}@example.com", password)
    assert client.get(f"/api/cases/{managed_case_id}", headers=ordinary_headers).status_code == 403
    assert managed_number not in {row["case_number"] for row in
                                  client.get("/api/cases", headers=ordinary_headers).json()}

    with SessionLocal() as db:
        db.execute(delete(Case).where(Case.case_number.in_([managed_number, other_number])))
        db.commit()
